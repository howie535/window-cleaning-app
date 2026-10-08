#!/usr/bin/env python3
"""Convert the Splash Out Round workbook into a single JSON file the app can import.

One-off migration tool for moving the round from the spreadsheet into the app.
It can be re-run as often as needed before the switchover; each run produces a
complete snapshot. Customer IDs are stable between runs (derived from round,
name and address), so a later import can replace an earlier one.

Source: "Splash Out Round.xlsx" (the single-sheet running workbook built from the
yearly worksheets). Only the typed-in cells are read; calculated columns are ignored.

Usage:
    python3 convert_round_workbook.py "Splash Out Round.xlsx" OUTPUT_FOLDER

Writes:
    OUTPUT_FOLDER/splash-out-import.json   the data (see Docs/SPEC.md, "Import file")
    OUTPUT_FOLDER/import-report.txt        counts, totals to check against the workbook, warnings

Never write the output inside the app's git repository: it holds customer data.

Needs: Python 3.9+ and openpyxl (pip install openpyxl).
"""
import datetime as dt
import json
import re
import sys
import uuid
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl

SCHEMA_VERSION = 1
ID_NAMESPACE = uuid.UUID("5b1e3f0a-6c2d-4f4e-9a51-53504c415348")  # fixed so IDs are stable

# Round sheet layout (see build_round.py)
COL = dict(round=1, pay=2, ring=3, front=4, since=5, notes=6, status=7, area=8, name=9,
           addr=10, price=11, eo=13)
FIRST_CLEAN_COL = 17     # column Q = "Clean 1"; payments follow in the next column
PAIRS = 120
DASHES = {"—", "–", "-", "— ", "--"}
EARLIEST_REAL_DATE = dt.date(2000, 1, 1)


def tax_year(d: dt.date) -> str:
    y = d.year if d >= dt.date(d.year, 4, 6) else d.year - 1
    return f"{y}-{str(y + 1)[-2:]}"


def as_date(v):
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    return None


def money(v):
    return round(float(v), 2) if isinstance(v, (int, float)) else None


def comment_text(cell):
    if not cell.comment:
        return None
    t = re.sub(r"\s+", " ", cell.comment.text).strip()
    return t.split("Comment:")[-1].strip() or None


def tick(v):
    return str(v or "").strip() == "✓"


def customer_id(round_name, name, address):
    key = "|".join(s.strip().lower() for s in (round_name, name, address))
    return str(uuid.uuid5(ID_NAMESPACE, key))


def tidy_dashes(visits, eo_flag, cancelled):
    """Give dashes their real meaning (in place).

    - Dashes before a customer's first clean: they hadn't started yet. Removed.
    - Dashes after the last clean of a cancelled customer: they'd left. Removed.
    - Every-other-clean customers: one dash between cleans is their normal off-cycle
      ("not_due"); in a longer run, every second dash is a real skip.
    - Anything else stays "skipped" (a one-off cancellation).
    """
    cleaned = [i for i, v in enumerate(visits) if v["kind"] == "cleaned"]
    if not cleaned:
        visits[:] = []
        return
    first, last = cleaned[0], cleaned[-1]
    keep = []
    run = []
    def flush(run):
        for j, v in enumerate(run):
            if eo_flag:
                v["kind"] = "skipped" if (j % 2 == 1) else "not_due"
            keep.append(v)
    for i, v in enumerate(visits):
        if v["kind"] == "skipped":
            if i < first or (cancelled and i > last):
                continue
            run.append(v)
        else:
            flush(run); run = []
            keep.append(v)
    flush(run)
    visits[:] = keep


def read_customers(ws, warnings):
    customers = []
    seq = 0
    for r in range(2, ws.max_row + 1):
        name = ws.cell(r, COL["name"]).value
        addr = ws.cell(r, COL["addr"]).value
        price = ws.cell(r, COL["price"]).value
        if not (name or addr):
            continue
        seq += 1
        name = str(name or "").strip()
        addr = str(addr or "").strip()
        rnd = str(ws.cell(r, COL["round"]).value or "").strip()
        notes = []
        for k in ("notes",):
            v = ws.cell(r, COL[k]).value
            if v:
                notes.append(str(v).strip())
        for k in ("name", "addr", "price", "eo"):
            t = comment_text(ws.cell(r, COL[k]))
            if t:
                notes.append(t)
        list_price = money(price) or 0.0
        visits = []
        last_date = None
        pending_skips = []   # skipped slots waiting for the next real clean to date them

        for k in range(PAIRS):
            dcell = ws.cell(r, FIRST_CLEAN_COL + 2 * k)
            pcell = ws.cell(r, FIRST_CLEAN_COL + 2 * k + 1)
            dv, pv = dcell.value, pcell.value
            if dv is None and pv is None and not dcell.comment and not pcell.comment:
                continue
            note = "; ".join(t for t in (comment_text(dcell), comment_text(pcell)) if t) or None
            d = as_date(dv)
            if d is not None:
                if d < EARLIEST_REAL_DATE:
                    warnings.append(f"{name} ({rnd}): clean {k + 1} has date {d} - not a real date, left out")
                    continue
                paid = money(pv)
                if paid is None and isinstance(pv, str) and pv.strip() not in DASHES and pv.strip():
                    note = "; ".join(x for x in (note, f"payment cell: {pv.strip()}") if x)
                paid = paid or 0.0
                charged = paid if paid > 0 else list_price
                visits.append(dict(date=d.isoformat(), kind="cleaned", slot=k + 1, list_price=list_price,
                                   charged=charged, paid=paid, note=note))
                for s in pending_skips:          # date earlier skips between the last clean and this one
                    s["date"] = (last_date + (d - last_date) / 2).isoformat() if last_date else d.isoformat()
                    s["date_estimated"] = True
                pending_skips = []
                last_date = d
            elif isinstance(dv, str) and dv.strip() in DASHES:
                s = dict(date=None, kind="skipped", slot=k + 1, list_price=list_price, charged=0.0, paid=0.0,
                         note=note)
                visits.append(s)
                pending_skips.append(s)
            elif isinstance(dv, str) and dv.strip():
                notes.append(f"Clean {k + 1}: {dv.strip()}" + (f" / {pv}" if pv not in (None, "") else ""))
            elif note:
                notes.append(f"Clean {k + 1}: {note}")
        for s in pending_skips:                  # skips after the last clean: one cycle on (6 weeks) each
            last_date = (last_date or dt.date.today()) + dt.timedelta(weeks=6)
            s["date"] = last_date.isoformat()
            s["date_estimated"] = True

        tidy_dashes(visits, eo_flag=str(ws.cell(r, COL["eo"]).value or "").strip().upper() == "EO",
                    cancelled=str(ws.cell(r, COL["status"]).value or "").strip() == "Cancelled")
        since = as_date(ws.cell(r, COL["since"]).value)
        contact = "ring" if tick(ws.cell(r, COL["ring"]).value) else "whatsapp"
        pay = str(ws.cell(r, COL["pay"]).value or "").strip().lower() or None
        customers.append(dict(
            id=customer_id(rnd, name, addr), sequence=seq, round=rnd,
            area=str(ws.cell(r, COL["area"]).value or "").strip() or None,
            status=str(ws.cell(r, COL["status"]).value or "Active").strip().lower().replace(" ", "_"),
            name=name, address=addr, phone=None, price=list_price,
            price_since=since.isoformat() if since else None,
            every_other=str(ws.cell(r, COL["eo"]).value or "").strip().upper() == "EO",
            front_only=tick(ws.cell(r, COL["front"]).value), contact=contact, pay_method=pay,
            notes=notes, visits=visits))
    return customers


def read_settings(wb):
    st = wb["Settings"]
    crews = []
    for r in range(4, 12):
        name, key, target = st.cell(r, 6).value, st.cell(r, 7).value, st.cell(r, 8).value
        if name and key:
            crews.append(dict(name=str(name), members=[{"H": "Howie", "D": "Dad", "R": "Rebekah"}[c] for c in str(key)],
                              day_target=money(target)))
    days = ["monday", "tuesday", "wednesday", "thursday", "friday", "saturday", "sunday"]
    usual = {}
    for n, day in enumerate(days):
        crew = st.cell(14 + n, 7).value
        usual[day] = None if (not crew or crew == "Not working") else str(crew)
    rise = as_date(st.cell(11, 2).value)
    return dict(
        crews=crews, usual_week=usual,
        extra_day_crew=str(st.cell(21, 7).value or "Howie + Dad"),
        min_houses_for_working_day=int(st.cell(22, 7).value or 5),
        overbook=float(st.cell(7, 2).value or 0.10),
        next_up_hide_weeks=int(st.cell(8, 2).value or 3),
        next_up_hide_weeks_every_other=int(st.cell(9, 2).value or 8),
        price_rise=dict(date=rise.isoformat() if rise else None, percent=float(st.cell(12, 2).value or 0.10),
                        rounding="nearest_pound_half_down", delay_months=int(st.cell(14, 2).value or 12),
                        due_after_months=int(st.cell(15, 2).value or 24)),
        records_start=(as_date(st.cell(24, 2).value) or dt.date(2026, 4, 6)).isoformat(),
    )


def read_diary(wb):
    out = []
    ws = wb["Diary"]
    for r in range(2, ws.max_row + 1):
        d = as_date(ws.cell(r, 1).value)
        if not d:
            continue
        members = [p for p, c in (("Howie", 3), ("Dad", 4), ("Rebekah", 5)) if tick(ws.cell(r, c).value)]
        out.append(dict(date=d.isoformat(), day_off=str(ws.cell(r, 2).value or "").strip() == "Day off",
                        crew=members or None, note=str(ws.cell(r, 6).value or "").strip() or None))
    return out


def read_tips(wb):
    out = []
    ws = wb["Tips"]
    for r in range(2, ws.max_row + 1):
        d, n, a = as_date(ws.cell(r, 1).value), ws.cell(r, 2).value, money(ws.cell(r, 3).value)
        if d and a:
            out.append(dict(date=d.isoformat(), name=str(n or "").strip(), amount=a))
    return out


def main(src, out_dir):
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.load_workbook(src)
    warnings = []
    customers = read_customers(wb["Round"], warnings)
    data = dict(schema=SCHEMA_VERSION, generated_at=dt.datetime.now().isoformat(timespec="seconds"),
                source=Path(src).name, settings=read_settings(wb), customers=customers,
                diary=read_diary(wb), tips=read_tips(wb))
    (out_dir / "splash-out-import.json").write_text(json.dumps(data, indent=1, ensure_ascii=False))

    # report: figures to tick off against the workbook before trusting the import
    status = Counter(c["status"] for c in customers)
    work_by_ty, cleans_by_ty, paid_by_ty = defaultdict(float), Counter(), defaultdict(float)
    owed, owed_n, skips, not_due = 0.0, 0, 0, 0
    for c in customers:
        for v in c["visits"]:
            if v["kind"] != "cleaned":
                skips += v["kind"] == "skipped"
                not_due += v["kind"] == "not_due"
                continue
            ty = tax_year(dt.date.fromisoformat(v["date"]))
            work_by_ty[ty] += v["charged"]
            cleans_by_ty[ty] += 1
            paid_by_ty[ty] += v["paid"]
            if v["paid"] == 0:
                owed += v["charged"]; owed_n += 1
    round_value = sum(c["price"] * (0.5 if c["every_other"] else 1) for c in customers
                      if c["status"] in ("active", "leaving", "not_started"))
    lines = [f"Splash Out import report  {data['generated_at']}", f"Source: {data['source']}", "",
             f"Customers: {len(customers)}  " + ", ".join(f"{k} {v}" for k, v in sorted(status.items())),
             f"Visits: {sum(cleans_by_ty.values())} cleans, {skips} one-off skips, {not_due} every-other off-cycles (their dates are estimated)",
             f"Diary entries: {len(data['diary'])}   Tips: {len(data['tips'])}", "",
             "Check these against the workbook (Today and Tax Year sheets):",
             f"  Round value per cycle: £{round_value:,.2f}",
             f"  Money owed: £{owed:,.2f} over {owed_n} cleans"]
    for ty in sorted(work_by_ty):
        lines.append(f"  Tax year {ty}: work £{work_by_ty[ty]:,.2f}, paid £{paid_by_ty[ty]:,.2f}, {cleans_by_ty[ty]} cleans")
    lines += ["", f"Warnings ({len(warnings)}):"] + [f"  {w}" for w in warnings]
    (out_dir / "import-report.txt").write_text("\n".join(lines) + "\n")
    print("\n".join(lines))


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    main(sys.argv[1], sys.argv[2])
